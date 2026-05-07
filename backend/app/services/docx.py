from io import BytesIO

from docx import Document

from app.schemas.meeting import DocxRequest


def build_docx_bytes(result: DocxRequest) -> bytes:
    doc = Document()
    doc.add_heading('Meeting Summary', level=1)

    doc.add_heading('Summary', level=2)
    doc.add_paragraph(result.summary or '')

    doc.add_heading('Participants', level=2)
    if result.participants:
        for p in result.participants:
            doc.add_paragraph(p, style='List Bullet')
    else:
        doc.add_paragraph('—')

    doc.add_heading('Decisions', level=2)
    if result.decisions:
        for d in result.decisions:
            doc.add_paragraph(d, style='List Bullet')
    else:
        doc.add_paragraph('—')

    doc.add_heading('Action Items', level=2)
    if result.action_items:
        for item in result.action_items:
            owner = f' ({item.owner})' if item.owner else ''
            due = f' [due: {item.due}]' if item.due else ''
            doc.add_paragraph(f'{item.task}{owner}{due}', style='List Bullet')
    else:
        doc.add_paragraph('—')

    doc.add_heading('Transcript', level=2)
    doc.add_paragraph(result.transcript or '')

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()

